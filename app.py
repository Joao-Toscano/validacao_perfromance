"""
AVD — Configuração e Geração de Bases
"""
import io, uuid, json, zipfile
import numpy as np
import pandas as pd
import streamlit as st
from openpyxl import load_workbook
from difflib import SequenceMatcher
 
st.set_page_config(page_title="AVD — Configuração", page_icon="📋",
                   layout="wide", initial_sidebar_state="collapsed")
 
st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700&display=swap');
html,body,[class*="css"]{ font-family:'Inter',sans-serif; }
.block-container{ padding:2rem 3rem 4rem !important; max-width:1100px; }
 
.steps-nav{ display:flex; gap:0; border-radius:10px; overflow:hidden;
            border:1px solid rgba(255,255,255,.07); margin-bottom:2rem; }
.sn{ flex:1; padding:13px 6px; text-align:center; font-size:12px; font-weight:500;
     background:rgba(255,255,255,.03); color:rgba(255,255,255,.3);
     border-right:1px solid rgba(255,255,255,.07); }
.sn:last-child{ border-right:none; }
.sn.active{ background:rgba(129,140,248,.15); color:#818cf8; font-weight:700; }
.sn.done  { background:rgba(16,185,129,.07); color:#34d399; }
.sn-num{ display:inline-flex; width:19px; height:19px; border-radius:50%;
         border:1.5px solid currentColor; font-size:10px; align-items:center;
         justify-content:center; margin-right:5px; }
.sn.done .sn-num{ background:#34d399; color:#064e3b; border-color:#34d399; }
 
.sh{ font-size:11px; font-weight:600; letter-spacing:1px; text-transform:uppercase;
     color:rgba(255,255,255,.3); margin:1.4rem 0 .5rem; }
 
.esc-row{ display:flex; gap:4px; margin:8px 0; }
.esc-pt { flex:1; background:rgba(255,255,255,.06); border:1px solid rgba(255,255,255,.08);
           border-radius:6px; padding:8px 4px; text-align:center;
           font-size:10px; color:rgba(255,255,255,.5); line-height:1.4; }
.esc-n  { font-size:15px; font-weight:700; color:#fff; display:block; }
 
.badge{ font-size:10px; padding:2px 8px; border-radius:99px; font-weight:500; border:1px solid; margin-right:3px; }
.b-blue { color:#93c5fd; border-color:rgba(147,197,253,.3); background:rgba(147,197,253,.07); }
.b-green{ color:#86efac; border-color:rgba(134,239,172,.3); background:rgba(134,239,172,.07); }
.b-red  { color:#fca5a5; border-color:rgba(252,165,165,.3); background:rgba(252,165,165,.07); }
.b-gray { color:rgba(255,255,255,.35); border-color:rgba(255,255,255,.1); background:rgba(255,255,255,.03); }
 
.logbox{ background:rgba(0,0,0,.4); border:1px solid rgba(255,255,255,.08); border-radius:8px;
         padding:14px 16px; font-family:monospace; font-size:12px; color:rgba(255,255,255,.7);
         max-height:360px; overflow-y:auto; line-height:1.9; }
.log-ok { color:#34d399; } .log-err{ color:#f87171; }
.log-inf{ color:#93c5fd; } .log-wrn{ color:#fbbf24; }
 
div[data-testid="stMetric"]{ background:rgba(255,255,255,.04);
  border:1px solid rgba(255,255,255,.08); border-radius:10px; padding:14px 18px; }
div[data-testid="stMetricValue"]{ font-size:1.7rem !important; font-weight:700; }
 
/* Tabela de perguntas */
.perg-table{ width:100%; border-collapse:collapse; font-size:12px; }
.perg-table th{ background:rgba(255,255,255,.06); color:rgba(255,255,255,.4);
               font-size:10px; letter-spacing:.5px; text-transform:uppercase;
               padding:8px 12px; text-align:left; border-bottom:1px solid rgba(255,255,255,.08); }
.perg-table td{ padding:8px 12px; border-bottom:1px solid rgba(255,255,255,.05);
               color:rgba(255,255,255,.75); vertical-align:top; }
.perg-table tr:hover td{ background:rgba(255,255,255,.02); }
.perg-num{ color:rgba(255,255,255,.25); font-family:monospace; }
.perg-txt{ max-width:300px; }
</style>
""", unsafe_allow_html=True)
 
# ── Constants ─────────────────────────────────────────────────────────────────
# Mapeamento exato dos cabeçalhos do arquivo do cliente
COL_MAP = {
    "question_set":      ["nome do bloco de perguntas","nome do bloco","bloco","question_set","block","grupo","seção","categoria"],
    "desc_question_set": ["descrição do bloco de perguntas","descrição do bloco","desc_question_set"],
    "competencia":       ["nome da competência","competência","competencia","skill","tema"],
    "pergunta":          ["pergunta","question","enunciado","texto","item"],
    "opcional":          ["essa pergunta é opcional","opcional","optional","essa pergunta é opcional?"],
    "aberta":            ["é uma pergunta aberta","aberta","open","tipo","é uma pergunta aberta?"],
    "escala":            ["qual o número do modelo da escala","qual o número do modelo da escala? (pergunta fechada)","escala","scale","número da escala"],
    "min_caracters":     ["número mínimo de caracteres","número mínimo de caracteres (opcional)","min caracters","min_caracters"],
    "max_caracters":     ["número máximo de caracteres","número máximo de caracteres (pergunta aberta)","max caracters","max_caracters"],
    "definicao":         ["descrição geral sobre a pergunta","definição","definicao","definition"],
    # Grupos avaliativos — colunas booleanas separadas
    "g_auto":            ["autoavaliação","autoavaliacao","auto"],
    "g_lider":           ["líder avaliando liderados","lider avaliando liderados","time","gestor avaliando"],
    "g_liderado":        ["liderados avaliando líder","liderados avaliando lider","gestor","liderado avaliando"],
    "g_par":             ["pares","par"],
    "g_stakeholder":     ["stakeholders","stakeholder"],
}
DESC_SCALE_SYNS = ["descrição do primeiro","descrição do segundo","descrição do terceiro",
                   "descrição do quarto","descrição do quinto","duplique"]
GRUPOS_OPTS = ["Autoavaliação","Time","Gestor","Par","Stakeholder"]
G_FIELDS = ["g_auto","g_lider","g_liderado","g_par","g_stakeholder"]
DICT_CORES = {
    2:['#A8002E','#006499'], 3:['#A8002E','#FFC252','#006499'],
    4:['#A8002E','#FFC252','#35AEFF','#033D17'], 5:['#A8002E','#FFC252','#35AEFF','#006499','#033D17'],
    6:['#A8002E','#FA6F26','#FFC252','#35AEFF','#006499','#033D17'],
}
OUTPUT_COLS = ["question_set","desc_question_set","competencia","pergunta",
               "opcional","aberta","escala","min caracters","max caracters","definição"]
 
# ── Helpers ───────────────────────────────────────────────────────────────────
def uid(): return str(uuid.uuid4())[:8]
 
def norm_bool(val):
    v = str(val or "").strip().lower()
    return "sim" if v in ("sim","yes","true","1","s","y") else "não"
 
def detect_col(headers_lower, field):
    syns = COL_MAP.get(field, [])
    for i, h in enumerate(headers_lower):
        h_clean = h.strip().rstrip("?").strip()
        if any(s in h_clean or h_clean in s for s in syns):
            return i
    return None
 
def sim_r(a, b):
    return SequenceMatcher(None, a.lower().strip(), b.lower().strip()).ratio()
 
def make_p(**kw):
    return {
        "id": uid(), "texto": kw.get("texto",""), "competencia": kw.get("competencia",""),
        "aberta": kw.get("aberta","não"), "escala": str(kw.get("escala","")),
        "opcional": kw.get("opcional","não"),
        "min_caracters": str(kw.get("min_c","")), "max_caracters": str(kw.get("max_c","")),
        "definicao": kw.get("definicao",""),
        "grupos": kw.get("grupos", list(GRUPOS_OPTS)),
    }
 
def make_b(nome="", desc="", perguntas=None, origem="manual"):
    return {"id": uid(), "nome": nome, "desc": desc, "perguntas": perguntas or [], "origem": origem}
 
def make_e(num=None, pontos=None):
    if num is None:
        num = max((e["num"] for e in ss.escalas), default=0) + 1
    try: num = int(float(str(num)))
    except: pass
    return {"id": uid(), "num": num, "pontos": pontos or ["", ""]}
 
# ── Session state ─────────────────────────────────────────────────────────────
DEFAULTS = {
    "step": 1, "escalas": [], "blocos": [], "bloco_sel": None,
    "form_perg": None,   # pergunta sendo adicionada/editada
    "cfg": {"cliente":"","token":"","id_rodada":"","id_qs":"","id_cq":"","id_oq":"",
            "val_min_nan":"0","autonomia_bp":True,"sso":True},
    "resultados": None, "log": []
}
for k, v in DEFAULTS.items():
    if k not in st.session_state: st.session_state[k] = v
ss = st.session_state
 
# ── Nav ───────────────────────────────────────────────────────────────────────
STEPS = ["1. Entrada", "2. Validação", "3. Configuração", "4. Geração"]
 
def steps_nav():
    pills = ""
    for i, label in enumerate(STEPS, 1):
        cls = "active" if i == ss.step else ("done" if i < ss.step else "")
        num = "✓" if i < ss.step else str(i)
        pills += f'<div class="sn {cls}"><span class="sn-num">{num}</span>{label}</div>'
    st.markdown(f'<div class="steps-nav">{pills}</div>', unsafe_allow_html=True)
 
def go(n): ss.step = n; st.rerun()
 
def nav_btns(back=None, fwd=None, fwd_label="Continuar →", fwd_disabled=False):
    c1, _, c2 = st.columns([1, 5, 1])
    if back and c1.button("← Voltar", use_container_width=True): go(back)
    if fwd  and c2.button(fwd_label, type="primary", use_container_width=True, disabled=fwd_disabled): go(fwd)
 
# ── Import xlsx ───────────────────────────────────────────────────────────────
def import_xlsx(file_bytes):
    wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    msgs = []
 
    # Escalas
    emap = {}
    for sn in wb.sheetnames:
        ws = wb[sn]
        data = list(ws.iter_rows(values_only=True))
        if not data: continue
        hl = [str(c or "").lower().strip() for c in data[0]]
        si = detect_col(hl, "escala")
        di = [i for i, h in enumerate(hl) if any(s in h for s in DESC_SCALE_SYNS)]
        for row in data[1:]:
            if si is None or si >= len(row): continue
            num = str(row[si] or "").strip()
            if not num: continue
            descs = [str(row[i] or "").strip() for i in di if i < len(row) and row[i]]
            if num not in emap: emap[num] = descs
            elif descs and not emap[num]: emap[num] = descs
    escalas_out = []
    for ns, descs in emap.items():
        try: n = int(float(ns))
        except: n = ns
        escalas_out.append(make_e(num=n, pontos=descs if descs else ["", ""]))
 
    # Blocos e perguntas
    bmap = {}
    for sn in wb.sheetnames:
        ws = wb[sn]
        data = list(ws.iter_rows(values_only=True))
        if len(data) < 2: continue
        raw_headers = [str(c or "").strip() for c in data[0]]
        hl = [h.lower() for h in raw_headers]
        idx = {f: detect_col(hl, f) for f in COL_MAP}
 
        def get(row, f):
            i = idx.get(f)
            return str(row[i] or "").strip() if (i is not None and i < len(row)) else ""
 
        def get_bool_grupo(row, f):
            i = idx.get(f)
            if i is None or i >= len(row): return True   # padrão: inclui
            return norm_bool(str(row[i] or "")) == "sim"
 
        for row in data[1:]:
            if not any(c not in (None, "") for c in row): continue
            bn  = get(row, "question_set")
            pt  = get(row, "pergunta")
            if not bn and not pt: continue
 
            # Grupos a partir das colunas booleanas
            grupos = []
            mapa_g = [("g_auto","Autoavaliação"),("g_lider","Time"),
                      ("g_liderado","Gestor"),("g_par","Par"),("g_stakeholder","Stakeholder")]
            for gf, gl in mapa_g:
                if get_bool_grupo(row, gf): grupos.append(gl)
            if not grupos: grupos = list(GRUPOS_OPTS)
 
            p = make_p(
                texto=pt, competencia=get(row,"competencia"),
                aberta=norm_bool(get(row,"aberta") or "não"),
                escala=get(row,"escala"),
                opcional=norm_bool(get(row,"opcional") or "não"),
                min_c=get(row,"min_caracters"), max_c=get(row,"max_caracters"),
                definicao=get(row,"definicao"), grupos=grupos,
            )
            bmap.setdefault(bn, {}).setdefault(sn, []).append(p)
 
    wb.close()
 
    blocos_out = []
    for bn, abas in bmap.items():
        al = list(abas.keys())
        if len(al) == 1:
            blocos_out.append(make_b(nome=bn, perguntas=abas[al[0]], origem="import"))
            continue
        sets = [set(p["texto"] for p in abas[a]) for a in al]
        if all(s == sets[0] for s in sets[1:]):
            blocos_out.append(make_b(nome=bn, perguntas=abas[al[0]], origem="import"))
            msgs.append(f"✓ **{bn}** idêntico em {len(al)} abas → mesclado.")
        else:
            for a in al:
                blocos_out.append(make_b(nome=f"{bn} — {a}", perguntas=abas[a], origem="import"))
            msgs.append(f"⚠ **{bn}** diferente em {len(al)} abas → renomeado com sufixo.")
 
    nomes = [b["nome"] for b in blocos_out]
    for i, n1 in enumerate(nomes):
        for n2 in nomes[i+1:]:
            if sim_r(n1, n2) > 0.85:
                msgs.append(f"⚠ Blocos parecidos: **{n1}** e **{n2}**")
 
    return blocos_out, escalas_out, msgs
 
# ── Validação ─────────────────────────────────────────────────────────────────
def validar():
    erros, avisos = [], []
    enum = {str(e["num"]) for e in ss.escalas}
    if not ss.escalas: erros.append("Nenhuma escala cadastrada.")
    if not ss.blocos:  erros.append("Nenhum bloco cadastrado.")
    seen = set()
    nomes = [b["nome"].strip() for b in ss.blocos]
    for n in nomes:
        if n in seen: erros.append(f'Bloco duplicado: "{n}"')
        seen.add(n)
    for i, n1 in enumerate(nomes):
        for n2 in nomes[i+1:]:
            if n1 != n2 and sim_r(n1, n2) > 0.85:
                avisos.append(f'Blocos muito parecidos: "{n1}" e "{n2}"')
    total = 0
    for b in ss.blocos:
        if not b["nome"].strip(): erros.append("Bloco sem nome.")
        if not b["perguntas"]:   avisos.append(f'Bloco "{b["nome"]}" sem perguntas.')
        for p in b["perguntas"]:
            total += 1
            if not p["texto"].strip(): erros.append(f'Pergunta vazia em "{b["nome"]}".')
            if p["aberta"] == "não":
                if not p["escala"]: avisos.append(f'Fechada sem escala em "{b["nome"]}".')
                elif p["escala"] not in enum: erros.append(f'Escala {p["escala"]} inexistente em "{b["nome"]}".')
            if p["aberta"] == "sim" and not p["max_caracters"]:
                avisos.append(f'Aberta sem máx chars em "{b["nome"]}".')
    return erros, avisos, total
 
# ── Script (portado do notebook) ──────────────────────────────────────────────
def run_script(cfg):
    log = []
    def L(msg, t="inf"): log.append((t, msg))
 
    try:
        id_rodada   = int(cfg["id_rodada"])
        id_qs       = int(cfg["id_qs"])
        id_cq       = int(cfg["id_cq"])
        id_oq       = int(cfg["id_oq"])
        val_min_nan = float(cfg.get("val_min_nan") or "0")
    except Exception as e:
        L(f"IDs inválidos: {e}", "err"); return log, {}
 
    # DataFrames internos
    rows_p = []
    for b in ss.blocos:
        for p in b["perguntas"]:
            rows_p.append({
                "question_set": b["nome"], "desc_question_set": b["desc"],
                "competencia": p["competencia"], "pergunta": p["texto"],
                "opcional": p["opcional"], "aberta": p["aberta"],
                "escala": int(float(p["escala"])) if p["aberta"]=="não" and p["escala"] else np.nan,
                "min caracters": int(p["min_caracters"]) if p["min_caracters"] else np.nan,
                "max caracters": int(p["max_caracters"]) if p["max_caracters"] else np.nan,
                "definição": p["definicao"] if p["definicao"] else np.nan,
            })
    base_perguntas = pd.DataFrame(rows_p)
 
    rows_e = []
    for e in ss.escalas:
        r = {"escala": e["num"],
             "min": float(e["pontos"][0]) if e["pontos"] and e["pontos"][0] else np.nan,
             "max": float(e["pontos"][-1]) if e["pontos"] and e["pontos"][-1] else np.nan}
        for i, pt in enumerate(e["pontos"], 1): r[f"descrição_{i}"] = pt if pt else np.nan
        rows_e.append(r)
    base_escalas = pd.DataFrame(rows_e)
 
    L(f"Perguntas: {len(base_perguntas)} | Escalas: {len(base_escalas)} linha(s)")
 
    # Question Sets
    L("Gerando Question Sets...")
    bqs = pd.DataFrame(columns=['name','description','id','order','evaluation_rounds'])
    first = True; cnt = 0; criados = []
    for _, row in base_perguntas.iterrows():
        if row['question_set'] not in criados:
            desc = "" if pd.isna(row.get('desc_question_set','')) else str(row.get('desc_question_set',''))
            reg  = {'id': id_qs if first else id_qs+cnt, 'order': cnt+1,
                    'name': row['question_set'], 'description': desc, 'evaluation_rounds': id_rodada}
            if first: first = False
            bqs.loc[len(bqs)] = reg; criados.append(row['question_set']); cnt += 1
    L(f"✓ {len(bqs)} question set(s)", "ok")
 
    # Open Questions
    L("Gerando Open Questions...")
    boq = pd.DataFrame(columns=['info','description','id','question_set','is_optional',
                                 'order','key','min_answer_length','max_answer_length','evaluation_rounds'])
    first = True; cnt_oq = 0; cnt_p = 1
    for _, row in base_perguntas.iterrows():
        if str(row['aberta']).lower() == 'não': cnt_p += 1; continue
        desc = ("" if pd.isna(row.get('competencia')) else str(row['competencia']).upper()+" | ")+row['pergunta']
        opcional = 0 if str(row['opcional']).lower() == 'não' else 1
        min_c = 0 if pd.isna(row.get('min caracters')) else int(row['min caracters'])
        max_c = int(row['max caracters'])
        qs = bqs[bqs['name'] == row['question_set']]
        reg = {'id': id_oq if first else id_oq+cnt_oq, 'order': cnt_p,
               'key': id_oq if first else id_oq+cnt_oq, 'description': desc, 'info': '',
               'is_optional': opcional, 'min_answer_length': min_c, 'max_answer_length': max_c,
               'question_set': qs['id'].iloc[0], 'evaluation_rounds': id_rodada}
        if first: first = False
        boq.loc[len(boq)] = reg; cnt_oq += 1; cnt_p += 1
    L(f"✓ {len(boq)} open question(s)", "ok")
 
    # Choice Questions
    L("Gerando Choice Questions...")
    bcq = pd.DataFrame(columns=['info','description','id','question_set','is_optional',
                                 'order','key','evaluation_rounds','escala'])
    first = True; cnt_cq = 0; cnt_p = 1
    for _, row in base_perguntas.iterrows():
        if str(row['aberta']).lower() == 'sim': cnt_p += 1; continue
        desc = ("" if pd.isna(row.get('competencia')) else str(row['competencia']).upper()+" | ")+row['pergunta']
        opcional = 0 if str(row['opcional']).lower() == 'não' else 1
        qs = bqs[bqs['name'] == row['question_set']]
        reg = {'id': id_cq if first else id_cq+cnt_cq, 'order': cnt_p,
               'key': id_cq if first else id_cq+cnt_cq, 'description': desc, 'info': '',
               'is_optional': opcional, 'question_set': qs['id'].iloc[0],
               'escala': int(row['escala']), 'evaluation_rounds': id_rodada}
        if first: first = False
        bcq.loc[len(bcq)] = reg; cnt_cq += 1; cnt_p += 1
    L(f"✓ {len(bcq)} choice question(s)", "ok")
 
    # Escalas
    L("Gerando itens de escala...")
    besc = pd.DataFrame(columns=['content','id','question','value','color'])
    for _, row in bcq.iterrows():
        bt = base_escalas[base_escalas['escala'] == row['escala']]
        if len(bt) == 0:
            L(f"Escala {row['escala']} não encontrada", "err"); continue
        for _, re_ in bt.iterrows():
            cols_d = [c for c in bt.columns if c not in ['escala','min','max']]
            pts = [c for c in cols_d if not pd.isna(re_.get(c))]
            cnt_pts = len(pts)
            minimo = re_.get('min'); maximo = float(re_['max'])
            flag_nan = pd.isna(minimo)
            if flag_nan:
                minimo = val_min_nan
                div = (maximo - float(minimo)) / max(cnt_pts-2, 1)
            else:
                minimo = float(minimo)
                div = (maximo - minimo) / max(cnt_pts-1, 1)
            cores = DICT_CORES.get(cnt_pts, ['#000000']*cnt_pts)
            ci = 0
            for col in cols_d:
                if pd.isna(re_.get(col)): continue
                if flag_nan:
                    reg = {'id':'','question':row['id'],'value':np.nan,'content':str(re_[col]),'color':'#000000'}
                    flag_nan = False
                else:
                    reg = {'id':'','question':row['id'],
                           'value': round(minimo + ci*div, 2),
                           'content': str(re_[col]),
                           'color': cores[ci] if ci < len(cores) else '#000000'}
                    ci += 1
                besc.loc[len(besc)] = reg
    L(f"✓ {len(besc)} item(ns) de escala", "ok")
 
    cliente = cfg["cliente"]
    resultados = {
        f"Question Sets {cliente}.json":    bqs.to_json(orient='records', force_ascii=False),
        f"Open Questions {cliente}.json":   boq.to_json(orient='records', force_ascii=False),
        f"Choice Questions {cliente}.json": bcq.drop(columns=['escala'],errors='ignore').to_json(orient='records', force_ascii=False),
        f"Escalas {cliente}.json":          besc.to_json(orient='records', force_ascii=False),
    }
    L("✓ JSONs prontos", "ok")
    return log, resultados
 
def build_zip(res):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf,'w',zipfile.ZIP_DEFLATED) as z:
        for fn, content in res.items():
            z.writestr(fn, content.encode('utf-8'))
    return buf.getvalue()
 
 
# ═══════════════════════════════════════════════════════════════════════════════
# STEP 1 — Entrada
# ═══════════════════════════════════════════════════════════════════════════════
if ss.step == 1:
    steps_nav()
    st.markdown("# Configuração de Avaliação")
 
    tab_imp, tab_man = st.tabs(["📥  Importar planilha", "✏️  Preencher manualmente"])
 
    # ── ABA IMPORT ────────────────────────────────────────────────────────────
    with tab_imp:
        st.caption("Arraste o arquivo .xlsx do cliente. O sistema detecta as colunas automaticamente.")
        up = st.file_uploader("Arquivo .xlsx", type=["xlsx","xls"], label_visibility="collapsed")
        if up:
            with st.spinner("Lendo arquivo..."):
                bi, ei, mi = import_xlsx(up.read())
            c1,c2,c3 = st.columns(3)
            c1.metric("Blocos", len(bi))
            c2.metric("Perguntas", sum(len(b["perguntas"]) for b in bi))
            c3.metric("Escalas", len(ei))
            if mi:
                with st.expander("📋 Ajustes automáticos realizados"):
                    for m in mi: st.markdown(m)
            ca, cb = st.columns(2)
            if ca.button("✓ Usar esses dados", type="primary", key="ir"):
                ss.blocos = bi; ss.escalas = ei; ss.bloco_sel = None; go(2)
            if ss.blocos and cb.button("＋ Adicionar aos dados existentes", key="ia"):
                ex = {b["nome"]: b for b in ss.blocos}
                for nb in bi:
                    if nb["nome"] in ex: ex[nb["nome"]]["perguntas"].extend(nb["perguntas"])
                    else: ss.blocos.append(nb)
                for ne in ei:
                    if not any(e["num"] == ne["num"] for e in ss.escalas): ss.escalas.append(ne)
                go(2)
 
    # ── ABA MANUAL ───────────────────────────────────────────────────────────
    with tab_man:
        escala_opts = [""] + [str(e["num"]) for e in ss.escalas]
 
        # ── Escalas ──
        st.markdown('<div class="sh">1. Escalas</div>', unsafe_allow_html=True)
        st.caption("Defina as escalas usadas nas perguntas fechadas.")
 
        for ei_idx, esc in enumerate(ss.escalas):
            with st.container(border=True):
                hc1, hc2 = st.columns([8,1])
                with hc1:
                    nv = st.number_input(f"Nº da escala", min_value=1,
                                         value=int(esc["num"]) if str(esc["num"]).replace(".","").isdigit() else ei_idx+1,
                                         key=f"en_{esc['id']}")
                    esc["num"] = nv
                    if any(str(p).strip() for p in esc["pontos"]):
                        prev = "".join(f'<div class="esc-pt"><span class="esc-n">{i+1}</span>{p or "—"}</div>'
                                       for i,p in enumerate(esc["pontos"]))
                        st.markdown(f'<div class="esc-row">{prev}</div>', unsafe_allow_html=True)
                    with st.expander("✏ Editar pontos"):
                        n_pts = st.number_input("Pontos", 2, 10, len(esc["pontos"]), key=f"np_{esc['id']}")
                        while len(esc["pontos"]) < n_pts: esc["pontos"].append("")
                        esc["pontos"] = esc["pontos"][:n_pts]
                        pc = st.columns(min(n_pts, 5))
                        for pi in range(n_pts):
                            esc["pontos"][pi] = pc[pi%len(pc)].text_input(
                                f"Ponto {pi+1}", value=esc["pontos"][pi], key=f"pt_{esc['id']}_{pi}",
                                placeholder=f"Ex: Nível {pi+1}")
                with hc2:
                    st.markdown("<br><br>", unsafe_allow_html=True)
                    if st.button("✕", key=f"de_{esc['id']}"): ss.escalas.pop(ei_idx); st.rerun()
 
        if st.button("＋ Adicionar escala", key="ae"): ss.escalas.append(make_e()); st.rerun()
        if not ss.escalas: st.info("Adicione ao menos uma escala para usar em perguntas fechadas.")
 
        st.divider()
 
        # ── Blocos ──
        st.markdown('<div class="sh">2. Blocos de perguntas</div>', unsafe_allow_html=True)
        col_bl, col_pr = st.columns([5, 7], gap="large")
 
        with col_bl:
            st.caption(f"{len(ss.blocos)} bloco(s) cadastrado(s)")
            for bi_idx, bloco in enumerate(ss.blocos):
                np_ = len(bloco["perguntas"])
                ne_ = sum(1 for p in bloco["perguntas"]
                          if not p["texto"] or (p["aberta"]=="não" and not p["escala"]))
                is_s = ss.bloco_sel == bloco["id"]
                with st.container(border=is_s):
                    r1, r2 = st.columns([5,1])
                    r1.markdown(f"**{bloco['nome'] or '(sem nome)'}**  \n"
                                f"<span style='font-size:11px;color:rgba(255,255,255,.4)'>{np_} pergunta(s)"
                                f"{' · '+str(ne_)+' erro(s)' if ne_ else ' · OK'}</span>",
                                unsafe_allow_html=True)
                    r2.markdown("")
                    b1, b2 = st.columns([3,1])
                    if b1.button("Editar", key=f"sb_{bloco['id']}", use_container_width=True):
                        ss.bloco_sel = bloco["id"] if not is_s else None; st.rerun()
                    if b2.button("✕", key=f"db_{bloco['id']}"):
                        ss.blocos = [b for b in ss.blocos if b["id"] != bloco["id"]]
                        if ss.bloco_sel == bloco["id"]: ss.bloco_sel = None
                        st.rerun()
 
            st.markdown("")
            if st.button("＋ Novo bloco", use_container_width=True, key="nb"):
                nb = make_b(nome="Novo bloco")
                ss.blocos.append(nb); ss.bloco_sel = nb["id"]; st.rerun()
 
        # ── Painel de perguntas ──
        with col_pr:
            bloco = next((b for b in ss.blocos if b["id"] == ss.bloco_sel), None)
 
            if bloco is None:
                st.info("👈 Selecione um bloco para editar suas perguntas.")
            else:
                bloco["nome"] = st.text_input("Nome do bloco *", value=bloco["nome"], key=f"bn_{bloco['id']}")
                bloco["desc"] = st.text_input("Descrição (opcional)", value=bloco["desc"],
                                               key=f"bd_{bloco['id']}", placeholder="Descrição do bloco...")
                st.markdown("")
 
                # ── Tabela de perguntas existentes ──
                if bloco["perguntas"]:
                    st.markdown(f"**{len(bloco['perguntas'])} pergunta(s) no bloco:**")
                    rows_tbl = []
                    for i, p in enumerate(bloco["perguntas"]):
                        rows_tbl.append({
                            "#": i+1,
                            "Pergunta": p["texto"][:70]+("…" if len(p["texto"])>70 else ""),
                            "Tipo": "Aberta" if p["aberta"]=="sim" else "Fechada",
                            "Escala": p["escala"] if p["aberta"]=="não" else "—",
                            "Opcional": "Sim" if p["opcional"]=="sim" else "Não",
                            "Grupos": ", ".join(p.get("grupos",[]))[:30],
                        })
                    df_tbl = pd.DataFrame(rows_tbl)
                    st.dataframe(df_tbl, use_container_width=True, hide_index=True,
                                 column_config={"#": st.column_config.NumberColumn(width="small")})
 
                    # Botões de ação por linha
                    cols_btns = st.columns(min(len(bloco["perguntas"]), 8))
                    for i, p in enumerate(bloco["perguntas"]):
                        with cols_btns[i % len(cols_btns)]:
                            if st.button(f"✕ #{i+1}", key=f"dp_{p['id']}", use_container_width=True,
                                         help=f"Remover pergunta {i+1}"):
                                bloco["perguntas"].pop(i); st.rerun()
 
                    st.divider()
 
                # ── Formulário para adicionar nova pergunta ──
                st.markdown("**Adicionar pergunta:**")
                with st.container(border=True):
                    novo_txt = st.text_area("Texto da pergunta *", key=f"new_txt_{bloco['id']}",
                                            placeholder="Digite o texto da pergunta...", height=80)
                    fa, fb = st.columns(2)
                    novo_comp = fa.text_input("Competência", key=f"new_comp_{bloco['id']}",
                                              placeholder="Ex: Trabalho em equipe")
                    novo_tipo = fb.selectbox("Tipo", ["Fechada","Aberta"], key=f"new_tipo_{bloco['id']}")
                    novo_aberta = "não" if novo_tipo == "Fechada" else "sim"
 
                    fc, fd = st.columns(2)
                    if novo_aberta == "não":
                        novo_escala = fc.selectbox("Escala *", escala_opts, key=f"new_esc_{bloco['id']}")
                        novo_min = ""; novo_max = ""
                    else:
                        novo_escala = ""
                        novo_min = fc.text_input("Mín. caracteres", key=f"new_min_{bloco['id']}", placeholder="0")
                        novo_max = fd.text_input("Máx. caracteres", key=f"new_max_{bloco['id']}", placeholder="10000")
 
                    novo_grupos = st.multiselect("Grupos avaliativos *", GRUPOS_OPTS,
                                                  default=GRUPOS_OPTS, key=f"new_grp_{bloco['id']}")
                    fe, ff = st.columns(2)
                    novo_opc = fe.checkbox("Pergunta opcional?", key=f"new_opc_{bloco['id']}")
                    novo_def = ff.text_input("Definição (opcional)", key=f"new_def_{bloco['id']}")
 
                    if st.button("＋ Adicionar pergunta ao bloco", type="primary",
                                 use_container_width=True, key=f"btn_add_{bloco['id']}"):
                        if not novo_txt.strip():
                            st.error("O texto da pergunta é obrigatório.")
                        elif novo_aberta == "não" and not novo_escala:
                            st.error("Selecione uma escala para pergunta fechada.")
                        else:
                            bloco["perguntas"].append(make_p(
                                texto=novo_txt, competencia=novo_comp, aberta=novo_aberta,
                                escala=novo_escala, opcional="sim" if novo_opc else "não",
                                min_c=novo_min, max_c=novo_max, definicao=novo_def,
                                grupos=novo_grupos,
                            ))
                            st.rerun()
 
    st.divider()
    erros, _, _ = validar()
    _, c_fwd = st.columns([5,1])
    with c_fwd:
        disabled = bool(erros) or not ss.blocos
        if st.button("Ir para Validação →", type="primary", use_container_width=True, disabled=disabled):
            go(2)
    if erros and ss.blocos:
        for e in erros[:3]: st.error(e)
 
 
# ═══════════════════════════════════════════════════════════════════════════════
# STEP 2 — Validação
# ═══════════════════════════════════════════════════════════════════════════════
elif ss.step == 2:
    steps_nav()
    st.markdown("# Visualização e Validação")
    erros, avisos, total_p = validar()
 
    m1,m2,m3,m4 = st.columns(4)
    m1.metric("Blocos", len(ss.blocos))
    m2.metric("Perguntas", total_p)
    m3.metric("Escalas", len(ss.escalas))
    m4.metric("Erros", len(erros), delta=str(len(erros)) if erros else None,
              delta_color="inverse" if erros else "off")
 
    st.divider()
    c_e, c_w = st.columns(2)
    with c_e:
        if erros:
            st.markdown("**❌ Erros (obrigatório corrigir)**")
            for e in erros: st.markdown(f"🔴 {e}")
        else:
            st.success("Sem erros!")
    with c_w:
        if avisos:
            st.markdown("**⚠️ Avisos**")
            for a in avisos: st.markdown(f"🟡 {a}")
        else:
            st.success("Sem avisos!")
 
    st.divider()
    st.markdown('<div class="sh">Prévia completa</div>', unsafe_allow_html=True)
    for b in ss.blocos:
        np_ = len(b["perguntas"])
        ne_ = sum(1 for p in b["perguntas"] if not p["texto"] or (p["aberta"]=="não" and not p["escala"]))
        with st.expander(f"{'🔴' if ne_ else '✅'} {b['nome']}  —  {np_} pergunta(s)", expanded=ne_>0):
            df_p = pd.DataFrame([{
                "#": i+1, "Pergunta": p["texto"],
                "Competência": p["competencia"],
                "Tipo": "Aberta" if p["aberta"]=="sim" else "Fechada",
                "Escala": p["escala"] if p["aberta"]=="não" else "—",
                "Opcional": "Sim" if p["opcional"]=="sim" else "Não",
                "Grupos": ", ".join(p.get("grupos",[])),
            } for i,p in enumerate(b["perguntas"])])
            st.dataframe(df_p, use_container_width=True, hide_index=True)
 
    st.divider()
    st.markdown('<div class="sh">Escalas</div>', unsafe_allow_html=True)
    for e in ss.escalas:
        with st.expander(f"Escala {e['num']} — {len(e['pontos'])} pontos"):
            prev = "".join(f'<div class="esc-pt"><span class="esc-n">{i+1}</span>{p or "—"}</div>'
                           for i,p in enumerate(e["pontos"]))
            st.markdown(f'<div class="esc-row">{prev}</div>', unsafe_allow_html=True)
 
    nav_btns(back=1, fwd=3, fwd_label="Configurar rodada →", fwd_disabled=bool(erros))
 
 
# ═══════════════════════════════════════════════════════════════════════════════
# STEP 3 — Configuração
# ═══════════════════════════════════════════════════════════════════════════════
elif ss.step == 3:
    steps_nav()
    st.markdown("# Configuração da Rodada")
    st.caption("Preencha as informações da rodada e os IDs gerados na plataforma antes de continuar.")
    cfg = ss.cfg
 
    st.markdown('<div class="sh">Acesso</div>', unsafe_allow_html=True)
    c1,c2 = st.columns(2)
    cfg["cliente"] = c1.text_input("Tenant do cliente *", value=cfg["cliente"], placeholder="ex: mindsight")
    cfg["token"]   = c2.text_input("Token do Performance *", value=cfg["token"], type="password")
 
    st.markdown('<div class="sh">IDs da rodada</div>', unsafe_allow_html=True)
    st.caption("Crie a rodada e os grupos na plataforma antes de preencher.")
    r1,r2 = st.columns(2)
    cfg["id_rodada"] = r1.text_input("ID da Rodada *", value=cfg["id_rodada"])
    cfg["id_qs"]     = r2.text_input("ID do 1º Question Set *", value=cfg["id_qs"])
    r3,r4 = st.columns(2)
    cfg["id_cq"] = r3.text_input("ID da 1ª Choice Question *", value=cfg["id_cq"])
    cfg["id_oq"] = r4.text_input("ID da 1ª Open Question *", value=cfg["id_oq"])
 
    st.markdown('<div class="sh">Configurações adicionais</div>', unsafe_allow_html=True)
    a1,a2,a3 = st.columns(3)
    cfg["val_min_nan"]  = a1.text_input("Valor mínimo p/ escala NaN", value=cfg["val_min_nan"])
    cfg["autonomia_bp"] = a2.checkbox("Autonomia de BPs?", value=cfg["autonomia_bp"])
    cfg["sso"]          = a3.checkbox("SSO?", value=cfg["sso"])
 
    ids_ok    = all(cfg[k].strip() for k in ["id_rodada","id_qs","id_cq","id_oq"])
    tenant_ok = bool(cfg["cliente"].strip())
    if not tenant_ok: st.warning("Preencha o tenant.")
    if not ids_ok:    st.warning("Preencha todos os IDs obrigatórios.")
 
    nav_btns(back=2, fwd=4, fwd_label="Gerar bases →", fwd_disabled=not(ids_ok and tenant_ok))
 
 
# ═══════════════════════════════════════════════════════════════════════════════
# STEP 4 — Geração
# ═══════════════════════════════════════════════════════════════════════════════
elif ss.step == 4:
    steps_nav()
    st.markdown("# Geração e Export")
 
    if ss.resultados is None:
        st.info("Clique no botão abaixo para processar os dados e gerar os arquivos JSON.")
        if st.button("▶  Gerar arquivos agora", type="primary", use_container_width=True):
            with st.spinner("Processando..."):
                log, res = run_script(ss.cfg)
            ss.log = log; ss.resultados = res; st.rerun()
    else:
        st.markdown('<div class="sh">Log de execução</div>', unsafe_allow_html=True)
        log_html = "".join(
            f'<div class="log-{t}">{"✓" if t=="ok" else "✕" if t=="err" else "→"} {msg}</div>'
            for t,msg in ss.log)
        st.markdown(f'<div class="logbox">{log_html}</div>', unsafe_allow_html=True)
 
        erros_exec = [m for t,m in ss.log if t=="err"]
        if erros_exec:
            st.error(f"{len(erros_exec)} erro(s). Volte e corrija os dados.")
        else:
            st.success("✅ Arquivos gerados com sucesso!")
            st.divider()
            st.markdown('<div class="sh">Downloads individuais</div>', unsafe_allow_html=True)
            for fname, content in ss.resultados.items():
                c1,c2 = st.columns([5,1])
                c1.markdown(f"**{fname}**  \n"
                            f"<span style='font-size:11px;color:rgba(255,255,255,.4)'>"
                            f"{len(json.loads(content))} registro(s)</span>", unsafe_allow_html=True)
                c2.download_button("⬇", data=content.encode("utf-8"), file_name=fname,
                                   mime="application/json", key=f"dl_{fname}", use_container_width=True)
            st.divider()
            st.download_button("⬇  Baixar todos (.zip)", data=build_zip(ss.resultados),
                               file_name=f"AVD_{ss.cfg['cliente']}.zip", mime="application/zip",
                               type="primary", use_container_width=True)
 
        st.divider()
        c1,c2 = st.columns([1,5])
        with c1:
            if st.button("← Voltar"): go(3)
        with c2:
            if st.button("↺ Regerar"): ss.resultados=None; ss.log=[]; st.rerun()
